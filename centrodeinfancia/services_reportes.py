"""Reporte XLSX de CDI, trabajadores y nómina acotado al alcance del usuario.

Las columnas y su orden replican el archivo validado por el equipo en el issue
#2508; la hoja de nómina agrega al final los indicadores RENAPER, calculados con
el mismo servicio que usa el PDF provincial.
"""

from __future__ import annotations

import json
import logging
from datetime import datetime
from io import BytesIO

from django.core.exceptions import FieldDoesNotExist
from django.db.models import Count, Exists, OuterRef
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from centrodeinfancia.access import (
    aplicar_scope_centros_cdi,
    tiene_alcance_simepi_nacional,
)
from centrodeinfancia.formulario_cdi_schema import (
    CAMPOS_OPCIONES,
    CAMPOS_OPCIONES_MULTIPLES,
    OPCIONES_DIAS_SEMANA,
)
from centrodeinfancia.models import (
    AccesoCDI,
    CentroDeInfancia,
    NominaCentroInfancia,
    Trabajador,
)
from centrodeinfancia.services_renaper_estado import (
    build_adult_validation_map,
    estado_renaper_nomina,
    motivo_renaper_nino,
)
from ciudadanos.models import Ciudadano
from core.models import Provincia


logger = logging.getLogger("django")

CHUNK_SIZE = 500
ORDEN_DIAS = {codigo: orden for orden, (codigo, _) in enumerate(OPCIONES_DIAS_SEMANA)}

COLUMNAS_CDI = (
    "id",
    "nombre",
    "codigo_cdi",
    "organizacion",
    "cuit_organizacion_gestiona",
    "provincia",
    "departamento",
    "decil_ipi",
    "nivel_inequidad_ipi",
    "municipio",
    "localidad",
    "codigo_postal",
    "ambito",
    "calle",
    "numero",
    "latitud",
    "longitud",
    "telefono",
    "mail",
    "nombre_referente",
    "apellido_referente",
    "email_referente",
    "telefono_referente",
    "dni_referente",
    "cuil_referente",
    "meses_funcionamiento",
    "dias_funcionamiento",
    "tipo_jornada",
    "tipo_jornada_otra",
    "modalidad_gestion",
    "modalidad_gestion_otra",
    "fecha_inicio",
    "fecha_creacion",
    "horarios_funcionamiento",
    "oferta_servicios",
    "referente_con_acceso_activo",
)

RUTAS_CDI = {
    "provincia": "provincia.nombre",
    "departamento": "departamento.nombre",
    "decil_ipi": "departamento.decil_ipi",
    "nivel_inequidad_ipi": "departamento.nivel_inequidad_ipi",
    "municipio": "municipio.nombre",
    "localidad": "localidad.nombre",
}

COLUMNAS_TRABAJADORES = (
    "id",
    "cdi_id",
    "cdi_nombre",
    "nombre",
    "apellido",
    "telefono",
    "rol",
    "fecha_carga",
    "subcomponente",
    "funcion_pfpi",
    "funcion_egp",
    "funcion_cdi",
    "sala_cdi",
    "funcion_uaf",
    "registro_tipo",
    "fecha_actualizacion",
    "fecha_nacimiento",
    "dni",
    "tipo_documentacion",
    "sexo_registral",
    "cuit",
    "pais_nacimiento",
    "nacionalidad_trabajador",
    "nivel_educativo",
    "formacion_academica",
    "capacitaciones_certificadas",
    "anos_trabajo_primera_infancia",
    "tipo_contratacion",
    "carga_horaria_semanal",
    "email",
    "calle_contacto",
    "unidad_funcional",
    "tipo_barrio",
    "provincia_contacto",
    "departamento_contacto",
    "municipio_contacto",
    "localidad_contacto",
    "grupo_pertenencia",
    "pueblo_originario",
    "lenguajes",
    "es_interprete",
    "tiene_discapacidad",
    "tipo_discapacidad",
    "recibe_apoyo_discapacidad",
    "tiene_cud",
    "numero_cud",
    "campos_verificados_renaper",
)

RUTAS_TRABAJADORES = {
    "cdi_id": "centro_id",
    "cdi_nombre": "centro.nombre",
    "provincia_contacto": "provincia_contacto.nombre",
    "municipio_contacto": "municipio_contacto.nombre",
    "localidad_contacto": "localidad_contacto.nombre",
}

COLUMNAS_NOMINA = (
    "id",
    "cdi_id",
    "cdi_nombre",
    "fecha",
    "estado",
    "dni",
    "apellido",
    "nombre",
    "fecha_nacimiento",
    "sexo",
    "nacionalidad",
    "sala",
    "pertenece_pueblo_originario",
    "pueblo_originario_cual",
    "habla_lengua_originaria_hogar",
    "talla",
    "peso",
    "calendario_vacunacion_al_dia",
    "tiene_discapacidad",
    "discapacidad_tipo",
    "recibe_apoyo_discapacidad",
    "posee_cud",
    "posee_obra_social",
    "calle_domicilio",
    "altura_domicilio",
    "piso_domicilio",
    "departamento_domicilio",
    "provincia_domicilio",
    "departamento_ipi_domicilio",
    "decil_ipi_domicilio",
    "municipio_domicilio",
    "localidad_domicilio",
    "responsable_legal_1_apellido",
    "responsable_legal_1_nombre",
    "responsable_legal_1_dni",
    "responsable_legal_1_telefono",
    "responsable_legal_1_percibe_auh",
    "responsable_legal_1_percibe_alimenta",
    "responsable_legal_1_relacion",
    "responsable_legal_1_fecha_nacimiento",
    "responsable_legal_1_tipo_documentacion",
    "responsable_legal_1_cuit",
    "responsable_legal_1_pais_nacimiento",
    "responsable_legal_1_nacionalidad",
    "responsable_legal_1_sexo_registral",
    "responsable_legal_1_nivel_educativo",
    "responsable_legal_1_consentimiento",
    "responsable_legal_2_apellido",
    "responsable_legal_2_nombre",
    "responsable_legal_2_dni",
    "responsable_legal_2_telefono",
    "responsable_legal_2_percibe_auh",
    "responsable_legal_2_percibe_alimenta",
    "responsable_legal_2_relacion",
    "responsable_legal_2_fecha_nacimiento",
    "responsable_legal_2_tipo_documentacion",
    "responsable_legal_2_cuit",
    "responsable_legal_2_pais_nacimiento",
    "responsable_legal_2_nacionalidad",
    "responsable_legal_2_sexo_registral",
    "responsable_legal_2_nivel_educativo",
    "responsable_legal_2_consentimiento",
    "adulto_responsable_apellido",
    "adulto_responsable_nombre",
    "adulto_responsable_dni",
    "adulto_responsable_telefono",
    "adulto_responsable_parentesco",
    "observaciones",
    "tipo_registro",
    "fecha_registro",
    "trabajador_registra_id",
    "tipo_documentacion",
    "cuit_nino",
    "pais_nacimiento",
    "edad_unidad",
    "tipo_barrio",
    "convivientes",
    "grupo_pertenencia",
    "lenguajes",
    "necesito_interprete",
    "numero_cud",
    "cobertura_salud",
    "controles_sanitarios_ultimo_anio",
    "longitud_acostado",
    "perimetro_cefalico",
    "lactancia",
    "diagnostico_peso",
    "diagnostico_talla",
    "orientacion_msal",
    "alergias_alimentarias",
    "anses_auh",
    "anses_aue",
    "anses_acsi",
    "anses_acn",
    "vacunacion_nomivac",
    "recibe_apoyo_desarrollo",
    # Agregadas por el issue #2508; no existen en el archivo original.
    "renaper_nino",
    "renaper_responsable_1",
    "renaper_responsable_2",
    "renaper_nino_motivo",
)

COLUMNAS_RESUMEN = ("seccion", "detalle", "cantidad")
COLUMNAS_DICCIONARIO = ("hoja", "columna", "codigo", "etiqueta")
COLUMNAS_METADATOS = ("dato", "valor")

RUTAS_NOMINA = {
    "cdi_id": "centro_id",
    "cdi_nombre": "centro.nombre",
    "provincia_domicilio": "provincia_domicilio.nombre",
    "departamento_ipi_domicilio": "departamento.nombre",
    "decil_ipi_domicilio": "departamento.decil_ipi",
    "municipio_domicilio": "municipio_domicilio.nombre",
    "localidad_domicilio": "localidad_domicilio.nombre",
}


def _resolver(obj, ruta):
    valor = obj
    for parte in ruta.split("."):
        valor = getattr(valor, parte, None)
        if valor is None:
            return None
    return valor


def _celda(valor):
    """Normaliza a un tipo que openpyxl acepte sin perder información."""
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, (list, dict)):
        return json.dumps(valor, ensure_ascii=False)
    if isinstance(valor, datetime) and timezone.is_aware(valor):
        # openpyxl rechaza datetimes con zona horaria.
        return timezone.localtime(valor).replace(tzinfo=None)
    return valor


def _valor(obj, columna, rutas, extra):
    if columna in extra:
        return extra[columna]
    return _resolver(obj, rutas.get(columna, columna))


def _fila(obj, columnas, rutas, extra=None):
    extra = extra or {}
    return [_celda(_valor(obj, columna, rutas, extra)) for columna in columnas]


def _horarios_funcionamiento(centro):
    horarios = sorted(
        centro.horarios_funcionamiento.all(),
        key=lambda horario: ORDEN_DIAS.get(horario.dia, len(ORDEN_DIAS)),
    )
    return (
        " | ".join(
            f"{horario.dia}: {horario.hora_apertura}-{horario.hora_cierre}"
            for horario in horarios
            if horario.hora_apertura and horario.hora_cierre
        )
        or None
    )


def _oferta_servicios(centro):
    # El orden viene del Meta.ordering de OfertaServicio (orden, codigo).
    return " | ".join(oferta.codigo for oferta in centro.oferta_servicios.all()) or None


def centros_en_alcance(user, provincia_id=None):
    """Única puerta de alcance del reporte: delega en el scope central de CDI.

    `provincia_id` solo acota; se aplica sobre el alcance ya resuelto, así que
    pedir una provincia ajena no agrega nada.
    """
    centros = aplicar_scope_centros_cdi(CentroDeInfancia.objects.all(), user)
    if str(provincia_id or "").isdecimal():
        centros = centros.filter(provincia_id=int(provincia_id))
    return centros


def provincias_en_alcance(user):
    """Provincias con centros visibles, para ofrecer el filtro sin revelar otras."""
    return Provincia.objects.filter(
        pk__in=centros_en_alcance(user).order_by().values("provincia_id")
    ).order_by("nombre")


def queryset_cdi(user, provincia_id=None):
    return (
        centros_en_alcance(user, provincia_id)
        .select_related("provincia", "departamento", "municipio", "localidad")
        .prefetch_related("horarios_funcionamiento", "oferta_servicios")
        .annotate(
            tiene_referente_activo=Exists(
                AccesoCDI.objects.filter(centro_id=OuterRef("pk"), activo=True)
            )
        )
        .order_by("nombre")
    )


def queryset_trabajadores(user, provincia_id=None):
    return (
        Trabajador.objects.filter(centro__in=centros_en_alcance(user, provincia_id))
        .select_related(
            "centro",
            "provincia_contacto",
            "municipio_contacto",
            "localidad_contacto",
        )
        .order_by("centro_id", "apellido", "nombre")
    )


def queryset_nomina(user, provincia_id=None):
    return (
        NominaCentroInfancia.objects.filter(
            centro__in=centros_en_alcance(user, provincia_id)
        )
        .select_related(
            "centro",
            "ciudadano",
            "provincia_domicilio",
            "departamento",
            "municipio_domicilio",
            "localidad_domicilio",
        )
        .order_by("centro_id", "apellido", "nombre")
    )


def _mapa_adultos(queryset):
    """Resuelve los indicadores de adultos en una consulta, no una por fila."""
    documentos = set()
    for documentos_fila in queryset.values_list(
        "responsable_legal_1_dni", "responsable_legal_2_dni"
    ):
        documentos.update(str(documento) for documento in documentos_fila if documento)
    return build_adult_validation_map(documentos)


def _abrir_hoja(workbook, titulo, columnas):
    hoja = workbook.create_sheet(titulo)
    # Deja fija la fila de encabezados al desplazarse.
    hoja.freeze_panes = "A2"
    hoja.append(list(columnas))
    return hoja


def _cerrar_hoja(hoja, columnas, filas):
    """Habilita el autofiltro sobre el rango realmente escrito."""
    ultima = get_column_letter(len(columnas))
    hoja.auto_filter.ref = f"A1:{ultima}{filas + 1}"
    return filas


def _contar_por(queryset, campo, etiquetas=None):
    """Agrupa contando en la base, sin recorrer las filas en Python."""
    conteos = queryset.values(campo).annotate(cantidad=Count("id")).order_by(campo)
    return [
        (
            (etiquetas or {}).get(fila[campo], fila[campo]) or "Sin dato",
            fila["cantidad"],
        )
        for fila in conteos
    ]


def _filas_resumen(user, provincia_id=None):
    centros = centros_en_alcance(user, provincia_id)
    trabajadores = queryset_trabajadores(user, provincia_id)
    nominas = queryset_nomina(user, provincia_id)

    filas = [
        ("Totales", "Centros de Desarrollo Infantil", centros.count()),
        ("Totales", "Trabajadores", trabajadores.count()),
        ("Totales", "Fichas de nómina", nominas.count()),
        (
            "Cobertura",
            "CDI sin fichas de nómina",
            centros.exclude(pk__in=nominas.order_by().values("centro_id")).count(),
        ),
        (
            "Cobertura",
            "CDI sin trabajadores",
            centros.exclude(pk__in=trabajadores.order_by().values("centro_id")).count(),
        ),
    ]
    for etiqueta, cantidad in _contar_por(
        nominas, "estado", dict(NominaCentroInfancia.ESTADO_CHOICES)
    ):
        filas.append(("Nómina por estado", etiqueta, cantidad))
    for etiqueta, cantidad in _contar_por(
        nominas,
        "ciudadano__estado_validacion_renaper",
        dict(Ciudadano.ESTADO_VALIDACION_RENAPER_CHOICES),
    ):
        filas.append(("Validación RENAPER de niños/as", etiqueta, cantidad))
    for etiqueta, cantidad in _contar_por(centros, "provincia__nombre"):
        filas.append(("CDI por provincia", etiqueta, cantidad))
    for etiqueta, cantidad in _contar_por(nominas, "centro__provincia__nombre"):
        filas.append(("Fichas de nómina por provincia", etiqueta, cantidad))
    return filas


_MODELOS_POR_HOJA = (
    ("CDI", CentroDeInfancia, COLUMNAS_CDI, RUTAS_CDI),
    ("Trabajadores", Trabajador, COLUMNAS_TRABAJADORES, RUTAS_TRABAJADORES),
    ("Nomina", NominaCentroInfancia, COLUMNAS_NOMINA, RUTAS_NOMINA),
)


def _opciones_de_columna(modelo, columna, rutas):
    """Etiquetas de una columna: primero el campo del modelo, luego el catálogo."""
    if columna not in rutas:
        try:
            campo = modelo._meta.get_field(columna)  # pylint: disable=protected-access
        except FieldDoesNotExist:
            campo = None
        if campo is not None and getattr(campo, "choices", None):
            return list(campo.choices)
    # Los campos JSON multivaluados no declaran choices; su catálogo vive aparte.
    return list(
        CAMPOS_OPCIONES.get(columna) or CAMPOS_OPCIONES_MULTIPLES.get(columna) or []
    )


def _filas_diccionario():
    filas = []
    for hoja, modelo, columnas, rutas in _MODELOS_POR_HOJA:
        for columna in columnas:
            for codigo, etiqueta in _opciones_de_columna(modelo, columna, rutas):
                filas.append((hoja, columna, codigo, etiqueta))
    return filas


def _filas_metadatos(user, totales, provincia_id=None):
    provincia = (
        Provincia.objects.filter(pk=provincia_id).first()
        if str(provincia_id or "").isdecimal()
        else None
    )
    alcance = (
        "Todos los centros"
        if tiene_alcance_simepi_nacional(user)
        else "Los centros del alcance del usuario"
    )
    return [
        ("Generado el", timezone.localtime().replace(tzinfo=None)),
        ("Generado por", getattr(user, "username", "") or "-"),
        ("Alcance", alcance),
        ("Filtro de provincia", provincia.nombre if provincia else "Sin filtro"),
        ("Filas hoja CDI", totales["cdi"]),
        ("Filas hoja Trabajadores", totales["trabajadores"]),
        ("Filas hoja Nomina", totales["nomina"]),
        (
            "Datos personales",
            "Contiene datos de niños, niñas y responsables. No compartir fuera "
            "del equipo.",
        ),
    ]


def _escribir_filas(workbook, titulo, columnas, filas):
    hoja = _abrir_hoja(workbook, titulo, columnas)
    for fila in filas:
        hoja.append(list(fila))
    return _cerrar_hoja(hoja, columnas, len(filas))


def _escribir_resumen(workbook, user, provincia_id=None):
    return _escribir_filas(
        workbook, "Resumen", COLUMNAS_RESUMEN, _filas_resumen(user, provincia_id)
    )


def _escribir_cdi(workbook, user, provincia_id=None):
    hoja = _abrir_hoja(workbook, "CDI", COLUMNAS_CDI)
    filas = 0
    for centro in queryset_cdi(user, provincia_id):
        extra = {
            "horarios_funcionamiento": _horarios_funcionamiento(centro),
            "oferta_servicios": _oferta_servicios(centro),
            "referente_con_acceso_activo": centro.tiene_referente_activo,
        }
        hoja.append(_fila(centro, COLUMNAS_CDI, RUTAS_CDI, extra))
        filas += 1
    return _cerrar_hoja(hoja, COLUMNAS_CDI, filas)


def _escribir_trabajadores(workbook, user, provincia_id=None):
    hoja = _abrir_hoja(workbook, "Trabajadores", COLUMNAS_TRABAJADORES)
    filas = 0
    for trabajador in queryset_trabajadores(user, provincia_id).iterator(
        chunk_size=CHUNK_SIZE
    ):
        hoja.append(_fila(trabajador, COLUMNAS_TRABAJADORES, RUTAS_TRABAJADORES))
        filas += 1
    return _cerrar_hoja(hoja, COLUMNAS_TRABAJADORES, filas)


def _escribir_nomina(workbook, user, provincia_id=None):
    hoja = _abrir_hoja(workbook, "Nomina", COLUMNAS_NOMINA)
    queryset = queryset_nomina(user, provincia_id)
    adultos = _mapa_adultos(queryset)
    filas = 0
    for registro in queryset.iterator(chunk_size=CHUNK_SIZE):
        extra = estado_renaper_nomina(registro, adult_validation=adultos)
        extra["renaper_nino_motivo"] = motivo_renaper_nino(registro)
        hoja.append(_fila(registro, COLUMNAS_NOMINA, RUTAS_NOMINA, extra))
        filas += 1
    return _cerrar_hoja(hoja, COLUMNAS_NOMINA, filas)


def generar_reporte_cdi_xlsx(user, provincia_id=None) -> bytes:
    """Devuelve el XLSX completo, solo con lo que el usuario puede ver."""
    workbook = Workbook(write_only=True)
    # El resumen va primero para que sea lo que se ve al abrir el archivo. Se
    # calcula con agregados, no recorriendo las filas de las otras hojas.
    _escribir_resumen(workbook, user, provincia_id)
    totales = {
        "cdi": _escribir_cdi(workbook, user, provincia_id),
        "trabajadores": _escribir_trabajadores(workbook, user, provincia_id),
        "nomina": _escribir_nomina(workbook, user, provincia_id),
    }
    _escribir_filas(workbook, "Diccionario", COLUMNAS_DICCIONARIO, _filas_diccionario())
    # Va último porque usa las filas realmente escritas, sin volver a contar.
    _escribir_filas(
        workbook,
        "Metadatos",
        COLUMNAS_METADATOS,
        _filas_metadatos(user, totales, provincia_id),
    )
    buffer = BytesIO()
    workbook.save(buffer)
    logger.info(
        "Reporte CDI generado",
        extra={
            "data": {
                "usuario_id": getattr(user, "pk", None),
                "provincia_id": provincia_id or None,
                **totales,
            }
        },
    )
    return buffer.getvalue()
