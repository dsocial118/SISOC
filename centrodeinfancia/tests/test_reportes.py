from datetime import date, time
from io import BytesIO

import openpyxl
import pytest
from django.contrib.auth.models import Group, Permission, User
from django.contrib.contenttypes.models import ContentType
from django.urls import reverse

from centrodeinfancia.models import (
    AccesoCDI,
    CentroDeInfancia,
    CentroDeInfanciaHorarioFuncionamiento,
    NominaCentroInfancia,
    OfertaServicio,
    Trabajador,
)
from centrodeinfancia.services_reportes import (
    COLUMNAS_CDI,
    COLUMNAS_NOMINA,
    COLUMNAS_TRABAJADORES,
    generar_reporte_cdi_xlsx,
)
from ciudadanos.models import Ciudadano
from core.constants import UserGroups
from core.models import Provincia
from users.bootstrap.groups_seed import permission_codes_for_bootstrap_group


# Contrato del archivo validado por el equipo en el issue #2508.
COLUMNAS_NOMINA_ORIGINALES = 96


def _usuario(username, *, superuser=False):
    return User.objects.create_user(
        username=username,
        password="test1234",
        is_superuser=superuser,
        is_staff=superuser,
    )


def _dar_permiso_vista(user):
    user.user_permissions.add(Permission.objects.get(codename="view_centrodeinfancia"))
    return User.objects.get(pk=user.pk)


def _dar_permiso_rol(user, codename, nombre):
    permiso, _ = Permission.objects.get_or_create(
        content_type=ContentType.objects.get_for_model(Group),
        codename=codename,
        defaults={"name": nombre},
    )
    user.user_permissions.add(permiso)
    return User.objects.get(pk=user.pk)


def _dar_permiso_exportacion(user):
    return _dar_permiso_rol(user, "role_exportar_a_csv", "Exportar a csv")


def _dar_permiso_reportes(user):
    return _dar_permiso_rol(user, "role_reportes_cdi", "Reportes CDI")


def _ciudadano(documento, *, validado=False, **extra):
    return Ciudadano.objects.create(
        apellido=extra.pop("apellido", "Perez"),
        nombre=extra.pop("nombre", "Ana"),
        fecha_nacimiento=date(2022, 3, 1),
        documento=documento,
        estado_validacion_renaper=(
            Ciudadano.RENAPER_VALIDADO if validado else Ciudadano.RENAPER_NO_CONSULTADO
        ),
        **extra,
    )


@pytest.fixture
def datos():
    provincia = Provincia.objects.create(nombre="Mendoza")
    propio = CentroDeInfancia.objects.create(nombre="CDI Propio", provincia=provincia)
    ajeno = CentroDeInfancia.objects.create(nombre="CDI Ajeno", provincia=provincia)

    Trabajador.objects.create(centro=propio, nombre="Ana", apellido="Propia")
    Trabajador.objects.create(centro=ajeno, nombre="Beto", apellido="Ajeno")

    NominaCentroInfancia.objects.create(
        centro=propio,
        ciudadano=_ciudadano(50000001),
        estado=NominaCentroInfancia.ESTADO_ACTIVO,
        apellido="Propia",
    )
    NominaCentroInfancia.objects.create(
        centro=ajeno,
        ciudadano=_ciudadano(50000002),
        estado=NominaCentroInfancia.ESTADO_ACTIVO,
        apellido="Ajena",
    )
    return {"propio": propio, "ajeno": ajeno}


def _hojas(contenido):
    libro = openpyxl.load_workbook(BytesIO(contenido))
    return {
        hoja.title: [list(fila) for fila in hoja.iter_rows(values_only=True)]
        for hoja in libro.worksheets
    }


def _columna(hoja, nombre):
    indice = hoja[0].index(nombre)
    return [fila[indice] for fila in hoja[1:]]


@pytest.mark.django_db
def test_encabezados_replican_el_contrato_del_issue(datos):
    hojas = _hojas(
        generar_reporte_cdi_xlsx(_usuario("reportes-encabezados", superuser=True))
    )

    assert list(hojas) == [
        "Resumen",
        "CDI",
        "Trabajadores",
        "Nomina",
        "Diccionario",
        "Metadatos",
    ]
    assert hojas["CDI"][0] == list(COLUMNAS_CDI)
    assert hojas["Trabajadores"][0] == list(COLUMNAS_TRABAJADORES)
    assert hojas["Nomina"][0] == list(COLUMNAS_NOMINA)
    assert len(COLUMNAS_CDI) == 36
    assert len(COLUMNAS_TRABAJADORES) == 47
    # Las columnas nuevas van al final para no mover las que el equipo ya validó.
    assert hojas["Nomina"][0][:COLUMNAS_NOMINA_ORIGINALES] == list(
        COLUMNAS_NOMINA[:COLUMNAS_NOMINA_ORIGINALES]
    )
    assert hojas["Nomina"][0][COLUMNAS_NOMINA_ORIGINALES:] == [
        "renaper_nino",
        "renaper_responsable_1",
        "renaper_responsable_2",
        "renaper_nino_motivo",
    ]


@pytest.mark.django_db
def test_superusuario_ve_todos_los_centros(datos):
    hojas = _hojas(generar_reporte_cdi_xlsx(_usuario("reportes-super", superuser=True)))

    assert _columna(hojas["CDI"], "nombre") == ["CDI Ajeno", "CDI Propio"]
    assert len(hojas["Trabajadores"]) == 3
    assert len(hojas["Nomina"]) == 3


@pytest.mark.django_db
def test_referente_solo_ve_su_centro_en_las_tres_hojas(datos):
    user = _usuario("reportes-referente")
    AccesoCDI.objects.create(user=user, centro=datos["propio"])

    hojas = _hojas(generar_reporte_cdi_xlsx(user))

    assert _columna(hojas["CDI"], "nombre") == ["CDI Propio"]
    assert _columna(hojas["Trabajadores"], "cdi_nombre") == ["CDI Propio"]
    assert _columna(hojas["Nomina"], "cdi_nombre") == ["CDI Propio"]


@pytest.mark.django_db
def test_usuario_sin_alcance_no_obtiene_filas(datos):
    user = _usuario("reportes-sin-alcance")
    grupo, _ = Group.objects.get_or_create(name=UserGroups.CDI_REFERENTE_CENTRO)
    user.groups.add(grupo)

    hojas = _hojas(generar_reporte_cdi_xlsx(User.objects.get(pk=user.pk)))

    assert hojas["CDI"] == [list(COLUMNAS_CDI)]
    assert hojas["Trabajadores"] == [list(COLUMNAS_TRABAJADORES)]
    assert hojas["Nomina"] == [list(COLUMNAS_NOMINA)]


@pytest.mark.django_db
def test_indicadores_renaper_de_nino_y_responsables():
    provincia = Provincia.objects.create(nombre="Mendoza")
    centro = CentroDeInfancia.objects.create(nombre="CDI RENAPER", provincia=provincia)
    _ciudadano(60000001, validado=True, apellido="Adulta", nombre="Rosa")
    NominaCentroInfancia.objects.create(
        centro=centro,
        ciudadano=_ciudadano(60000002, validado=True),
        estado=NominaCentroInfancia.ESTADO_ACTIVO,
        responsable_legal_1_dni=60000001,
        responsable_legal_2_dni=60000003,
    )

    hojas = _hojas(
        generar_reporte_cdi_xlsx(_usuario("reportes-renaper", superuser=True))
    )
    fila = hojas["Nomina"][1]
    encabezados = hojas["Nomina"][0]

    assert fila[encabezados.index("renaper_nino")] == "Sí"
    assert fila[encabezados.index("renaper_responsable_1")] == "Sí"
    # El responsable 2 no existe como Ciudadano: el indicador no puede afirmarlo.
    assert fila[encabezados.index("renaper_responsable_2")] == "No"


@pytest.mark.django_db
def test_motivo_explica_por_que_el_indicador_no_dice_si(datos):
    provincia = Provincia.objects.get(nombre="Mendoza")
    centro = CentroDeInfancia.objects.create(nombre="CDI Motivos", provincia=provincia)
    sin_consultar = _ciudadano(61000001)
    no_validado = _ciudadano(61000002)
    Ciudadano.objects.filter(pk=no_validado.pk).update(
        estado_validacion_renaper=Ciudadano.RENAPER_NO_VALIDADO,
        motivo_no_validacion_renaper=Ciudadano.MOTIVO_NO_VALIDADO_OTRO,
        motivo_no_validacion_descripcion="No coincide el apellido con RENAPER.",
    )
    for ciudadano, apellido in ((sin_consultar, "Aaa"), (no_validado, "Bbb")):
        NominaCentroInfancia.objects.create(
            centro=centro,
            ciudadano=ciudadano,
            estado=NominaCentroInfancia.ESTADO_ACTIVO,
            apellido=apellido,
        )

    hojas = _hojas(
        generar_reporte_cdi_xlsx(_usuario("reportes-motivo", superuser=True))
    )
    motivos = dict(
        zip(
            _columna(hojas["Nomina"], "apellido"),
            _columna(hojas["Nomina"], "renaper_nino_motivo"),
        )
    )

    # Distinguir "nunca se consultó" de "se consultó y no coincide" es lo accionable.
    assert motivos["Aaa"] == "No consultado"
    assert motivos["Bbb"] == "No coincide el apellido con RENAPER."


@pytest.mark.django_db
def test_resumen_es_la_primera_hoja_y_cuenta_lo_que_el_usuario_ve(datos):
    user = _usuario("reportes-resumen-referente")
    AccesoCDI.objects.create(user=user, centro=datos["propio"])

    hojas = _hojas(generar_reporte_cdi_xlsx(user))
    resumen = hojas["Resumen"]

    assert resumen[0] == ["seccion", "detalle", "cantidad"]
    totales = {
        detalle: cantidad
        for seccion, detalle, cantidad in resumen[1:]
        if seccion == "Totales"
    }
    # Solo su centro: 1 CDI, 1 trabajador y 1 ficha, no los del centro ajeno.
    assert totales == {
        "Centros de Desarrollo Infantil": 1,
        "Trabajadores": 1,
        "Fichas de nómina": 1,
    }
    secciones = {seccion for seccion, _, _ in resumen[1:]}
    assert "Validación RENAPER de niños/as" in secciones
    assert "CDI por provincia" in secciones


@pytest.mark.django_db
def test_hojas_con_encabezado_fijo_y_autofiltro(datos):
    contenido = generar_reporte_cdi_xlsx(_usuario("reportes-formato", superuser=True))
    libro = openpyxl.load_workbook(BytesIO(contenido))

    for hoja in libro.worksheets:
        assert hoja.freeze_panes == "A2", hoja.title
        assert hoja.auto_filter.ref is not None, hoja.title
        assert hoja.auto_filter.ref.startswith("A1:"), hoja.title


@pytest.mark.django_db
def test_diccionario_traduce_los_codigos_que_publican_las_hojas(datos):
    hojas = _hojas(generar_reporte_cdi_xlsx(_usuario("reportes-dicc", superuser=True)))
    diccionario = hojas["Diccionario"]

    assert diccionario[0] == ["hoja", "columna", "codigo", "etiqueta"]
    entradas = {(fila[0], fila[1], fila[2]): fila[3] for fila in diccionario[1:]}
    # Un código de cada hoja, para verificar que las tres quedan cubiertas.
    assert entradas[("CDI", "ambito", "urbano")] == "Urbano"
    assert entradas[("Nomina", "estado", "activo")] == "Activo"
    assert any(hoja == "Trabajadores" for hoja, _, _ in entradas)


@pytest.mark.django_db
def test_metadatos_declaran_origen_alcance_y_filas(datos):
    user = _usuario("reportes-metadatos-referente")
    AccesoCDI.objects.create(user=user, centro=datos["propio"])

    hojas = _hojas(generar_reporte_cdi_xlsx(user))
    metadatos = dict(hojas["Metadatos"][1:])

    assert metadatos["Generado por"] == "reportes-metadatos-referente"
    assert metadatos["Alcance"] == "Los centros del alcance del usuario"
    assert metadatos["Filtro de provincia"] == "Sin filtro"
    assert metadatos["Filas hoja CDI"] == 1
    assert metadatos["Filas hoja Nomina"] == 1
    assert "No compartir" in metadatos["Datos personales"]


@pytest.mark.django_db
def test_filtro_de_provincia_acota_sin_ampliar_el_alcance(client, datos):
    otra = Provincia.objects.create(nombre="San Juan")
    CentroDeInfancia.objects.create(nombre="CDI Otra Provincia", provincia=otra)
    mendoza = Provincia.objects.get(nombre="Mendoza")
    user = _dar_permiso_reportes(_dar_permiso_vista(_usuario("reportes-filtro")))
    AccesoCDI.objects.create(user=user, centro=datos["propio"])
    client.force_login(user)

    hojas = _hojas(generar_reporte_cdi_xlsx(user, mendoza.pk))
    assert _columna(hojas["CDI"], "nombre") == ["CDI Propio"]
    assert dict(hojas["Metadatos"][1:])["Filtro de provincia"] == "Mendoza"

    # Pedir una provincia fuera del alcance no agrega centros: acota, no amplía.
    respuesta = client.get(
        reverse("centrodeinfancia_reportes_descargar"), {"provincia": otra.pk}
    )
    assert respuesta.status_code == 200
    assert _columna(_hojas(respuesta.content)["CDI"], "nombre") == []


@pytest.mark.django_db
def test_pantalla_ofrece_solo_las_provincias_del_alcance(client, datos):
    Provincia.objects.create(nombre="San Juan")
    user = _dar_permiso_reportes(_dar_permiso_vista(_usuario("reportes-provincias")))
    AccesoCDI.objects.create(user=user, centro=datos["propio"])
    client.force_login(user)

    respuesta = client.get(reverse("centrodeinfancia_reportes"))

    nombres = [provincia.nombre for provincia in respuesta.context["provincias"]]
    assert nombres == ["Mendoza"]


@pytest.mark.django_db
def test_columnas_derivadas_del_centro():
    provincia = Provincia.objects.create(nombre="Mendoza")
    centro = CentroDeInfancia.objects.create(
        nombre="CDI Derivadas", provincia=provincia
    )
    # Se cargan desordenados para comprobar el orden de la semana.
    CentroDeInfanciaHorarioFuncionamiento.objects.create(
        centro=centro,
        dia="martes",
        hora_apertura=time(9, 0),
        hora_cierre=time(16, 0),
    )
    CentroDeInfanciaHorarioFuncionamiento.objects.create(
        centro=centro,
        dia="lunes",
        hora_apertura=time(8, 0),
        hora_cierre=time(15, 0),
    )
    centro.oferta_servicios.add(
        OfertaServicio.objects.create(codigo="tres_anos", orden=2),
        OfertaServicio.objects.create(codigo="dos_anos", orden=1),
    )
    usuario_referente = _usuario("reportes-derivadas-referente")
    AccesoCDI.objects.create(user=usuario_referente, centro=centro, activo=True)

    hojas = _hojas(
        generar_reporte_cdi_xlsx(_usuario("reportes-derivadas", superuser=True))
    )

    assert _columna(hojas["CDI"], "horarios_funcionamiento") == [
        "lunes: 08:00:00-15:00:00 | martes: 09:00:00-16:00:00"
    ]
    assert _columna(hojas["CDI"], "oferta_servicios") == ["dos_anos | tres_anos"]
    assert _columna(hojas["CDI"], "referente_con_acceso_activo") == [1]


@pytest.mark.django_db
def test_centro_sin_horarios_ni_oferta_no_rompe_el_reporte():
    provincia = Provincia.objects.create(nombre="Mendoza")
    CentroDeInfancia.objects.create(nombre="CDI Vacio", provincia=provincia)

    hojas = _hojas(generar_reporte_cdi_xlsx(_usuario("reportes-vacio", superuser=True)))

    assert _columna(hojas["CDI"], "horarios_funcionamiento") == [None]
    assert _columna(hojas["CDI"], "oferta_servicios") == [None]
    assert _columna(hojas["CDI"], "referente_con_acceso_activo") == [0]


@pytest.mark.django_db
@pytest.mark.parametrize(
    "otorgar", [_dar_permiso_reportes, _dar_permiso_exportacion], ids=["propio", "csv"]
)
def test_descarga_exige_permiso_de_reportes(client, datos, otorgar):
    user = _dar_permiso_vista(_usuario(f"reportes-descarga-{otorgar.__name__}"))
    client.force_login(user)

    respuesta = client.get(reverse("centrodeinfancia_reportes_descargar"))
    assert respuesta.status_code == 403

    client.force_login(otorgar(user))
    respuesta = client.get(reverse("centrodeinfancia_reportes_descargar"))

    assert respuesta.status_code == 200
    assert respuesta["Cache-Control"] == "private, no-store"
    assert "attachment" in respuesta["Content-Disposition"]
    assert "Resumen" in _hojas(respuesta.content)


@pytest.mark.parametrize(
    "grupo",
    [
        "SIMEPI - Administrador",
        "SIMEPI - Analista de datos",
        "SIMEPI - Equipo Nacional",
        "SIMEPI - Auditoría",
        "SIMEPI - EGP",
    ],
)
def test_roles_simepi_tienen_el_permiso_de_reportes(grupo):
    """Sin esto el módulo queda inaccesible para quienes lo pidieron."""
    assert "auth.role_reportes_cdi" in permission_codes_for_bootstrap_group(grupo)


@pytest.mark.django_db
def test_pantalla_avisa_cuando_no_puede_exportar(client, datos):
    user = _dar_permiso_vista(_usuario("reportes-pantalla"))
    client.force_login(user)

    respuesta = client.get(reverse("centrodeinfancia_reportes"))

    assert respuesta.status_code == 200
    assert respuesta.context["puede_exportar"] is False
